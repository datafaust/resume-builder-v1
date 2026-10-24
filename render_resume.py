#!/usr/bin/env python3
"""Render a resume from a YAML accomplishment inventory + JSON selection.
Usage: python render_resume.py --inventory master_accomplishments.yaml --selection selection.json --template resume-template.docx --output generated-resume.docx
Selection filenames are read from generated_selections beside this script.
Template filenames are read from templates beside this script.
The output file is saved in generated_resumes beside this script.
Requires: python-docx, pyyaml. LibreOffice can export the output to PDF.
"""
import argparse, json
from pathlib import Path
import yaml
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_bullet(doc, text):
    p = doc.add_paragraph(style='Resume Bullet')
    p.add_run('•  ' + text)
    return p

def add_section(doc, title):
    p = doc.add_paragraph(style='Resume Section')
    p.add_run(title.upper())
    return p

def add_two_side(doc, left, right, style, bold_left=True):
    # Explicit 2-cell row avoids fragile tabs and is well supported in LibreOffice.
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(5.65)
    table.columns[1].width = Inches(1.83)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcPr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement('w:tcMar')
        for side in ('top','left','bottom','right'):
            e=OxmlElement('w:'+side); e.set(qn('w:w'),'0'); e.set(qn('w:type'),'dxa');mar.append(e)
        tcPr.append(mar)
    a,b=table.rows[0].cells
    p=a.paragraphs[0];p.style=doc.styles[style];p.add_run(left).bold=bold_left
    p=b.paragraphs[0];p.style=doc.styles[style];p.alignment=WD_ALIGN_PARAGRAPH.RIGHT;p.add_run(right)
    return table



def add_summary(doc, inv, summary):
    """Render an AI-authored paragraph; source IDs are validated, never concatenated.

    Also accepts the legacy list of bullet strings for backward compatibility.
    """
    if isinstance(summary, list):
        if summary:
            add_section(doc, 'Summary')
            for item in summary:
                if not isinstance(item, str):
                    raise ValueError('Legacy summary must contain text strings')
                add_bullet(doc, item)
        return
    if not isinstance(summary, dict):
        raise ValueError('summary must be {"text": ..., "source_ids": [...]}')
    summary_text = summary.get('text', '')
    ids = summary.get('source_ids', [])
    if not isinstance(summary_text, str) or not summary_text.strip():
        raise ValueError('summary.text must be a non-empty paragraph')
    if not isinstance(ids, list) or not ids or not all(isinstance(i, str) for i in ids):
        raise ValueError('summary.source_ids must be a non-empty list of IDs')
    known_ids = {x['id'] for x in inv.get('summary_statements', [])}
    for group in ('experience', 'projects'):
        for entry in inv.get(group, []):
            known_ids.update(x['id'] for x in entry.get('accomplishments', []))
    unknown = sorted(set(ids) - known_ids)
    if unknown:
        raise ValueError('Unknown summary source IDs: ' + ', '.join(unknown))
    if '\n' in summary_text.strip():
        raise ValueError('summary.text must be a single paragraph (no newlines)')
    add_section(doc, 'Professional Summary')
    doc.add_paragraph(summary_text.strip(), style='Resume Summary')

def produce(inv, sel, template, output):
    template_dir = Path(__file__).resolve().parent / 'templates'
    template_dir.mkdir(parents=True, exist_ok=True)
    template_path = Path(template)
    if template_path.parent == Path('.'):
        template_path = template_dir / template_path
    if not template_path.is_file():
        raise FileNotFoundError(f'Resume template not found: {template_path}')
    doc=Document(template_path)
    # Template is a blank style carrier, not a sample with placeholders.
    body=doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'): body.remove(child)
    name = sel.get('name') or inv['identity']['name'].title()
    p=doc.add_paragraph(style='Resume Name');p.add_run(name)
    p=doc.add_paragraph(style='Resume Contact');p.add_run(sel.get('contact_line') or inv['identity']['contact_line_as_written'])
    if sel.get('summary'):
        add_summary(doc, inv, sel['summary'])
    if sel.get('skills'):
        add_section(doc,'Skills & Tools')
        for s in sel['skills']:
            p=doc.add_paragraph(style='Resume Skill Label');p.add_run(s['category']+':')
            p=doc.add_paragraph(style='Resume Skill Value');p.add_run(s['text'])
    if sel.get('certifications'):
        add_section(doc,'Certifications')
        for c in sel['certifications']: add_bullet(doc,c)
    if sel.get('experience'):
        add_section(doc,'Work Experience')
        roles={e['id']:e for e in inv['experience']}
        for wanted in sel['experience']:
            role=roles[wanted['id']]
            add_two_side(doc,role['employer'],role['location'],'Resume Company')
            add_two_side(doc,role['title'],wanted.get('dates') or role['dates_as_written'],'Resume Job',False)
            lookup = {
                a['id']: a['original_text']
                for a in role['accomplishments']
            }

            overrides = wanted.get('bullet_overrides', {})

            # Validate that rewrites refer only to selected accomplishments.
            invalid = set(overrides) - set(wanted['bullets'])
            if invalid:
                raise ValueError(f'Overrides reference unselected bullets: {invalid}')

            for aid in wanted['bullets']:
                if aid not in lookup:
                    raise ValueError(f'Unknown accomplishment ID: {aid}')

                text = overrides.get(aid, lookup[aid])

                if not isinstance(text, str) or not text.strip():
                    raise ValueError(f'Invalid bullet text: {aid}')

                add_bullet(doc, text.strip())
    if sel.get('projects'):
        add_section(doc,'Projects')
        roles={e['id']:e for e in inv['projects']}
        for wanted in sel['projects']:
            proj=roles[wanted['id']]
            p=doc.add_paragraph(style='Resume Project');p.add_run(proj['name'])
            p=doc.add_paragraph(style='Resume Job');p.add_run(proj['descriptor'])
            lookup={a['id']:a['original_text'] for a in proj['accomplishments']}
            for aid in wanted['bullets']: add_bullet(doc,lookup[aid])
    if sel.get('education'):
        add_section(doc,'Education')
        for e in inv['education']:
            add_two_side(doc,e['institution'],e['location'],'Resume Company')
            add_two_side(doc,e['degree_as_written'],e['date_as_written'],'Resume Job',False)
    if sel.get('publications'):
        add_section(doc,'Publications')
        pubs={p['id']:p['original_text'] for p in inv['publications']}
        for pid in sel['publications']:
            if pid not in pubs: raise ValueError('Unknown publication '+pid)
            add_bullet(doc,pubs[pid])
    output_dir = Path(__file__).resolve().parent / 'generated_resumes'
    output_dir.mkdir(parents=True, exist_ok=True)
    output_path = output_dir / Path(output).name
    doc.save(output_path)
    return output_path

if __name__=='__main__':
    ap=argparse.ArgumentParser();ap.add_argument('--inventory',default='master_accomplishments.yaml');ap.add_argument('--selection',default='selection.json');ap.add_argument('--template',default='resume-template.docx');ap.add_argument('--output',default='generated-resume.docx')
    x=ap.parse_args()
    selection_path = Path(x.selection)
    if selection_path.parent == Path('.'):
        selection_path = Path(__file__).resolve().parent / 'generated_selections' / selection_path
    with open(x.inventory,encoding='utf-8') as f: inventory=yaml.safe_load(f)
    with selection_path.open(encoding='utf-8') as f: selection=json.load(f)
    output_path = produce(inventory,selection,x.template,x.output)
    print('Wrote', output_path)
